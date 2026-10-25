"""Ingest security gates and the source-doc-id helper.

Two gates guard the only place untrusted input enters the system:
validate_local_path (path traversal / extension / size) and validate_url
(SSRF). Both must run before any content is read or fetched.
"""

import hashlib
import ipaddress
import socket
from pathlib import Path
from urllib.parse import urlsplit

import requests
from langchain_core.documents import Document

from research_assistant.config import get_settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
ALLOWED_URL_SCHEMES = {"http", "https"}
MAX_REDIRECTS = 3


class UnsupportedFileTypeError(ValueError):
    pass


class PathSecurityError(ValueError):
    pass


class UrlSecurityError(ValueError):
    pass


def make_source_doc_id(source: str) -> str:
    """Stable short id for a source path/URL; used for bulk-delete-by-document later."""
    return hashlib.sha256(source.encode()).hexdigest()[:12]


def validate_local_path(path: Path, base_dir: Path) -> Path:
    """Resolve `path` and ensure it stays inside `base_dir`, has an allowed
    extension, and is within the configured size limit."""
    settings = get_settings()
    resolved_base = base_dir.resolve()
    resolved = (base_dir / path).resolve() if not path.is_absolute() else path.resolve()

    if not resolved.is_relative_to(resolved_base):
        raise PathSecurityError(f"Path {path} escapes base directory {base_dir}")

    if resolved.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"Unsupported file extension: {resolved.suffix}")

    if not resolved.is_file():
        raise FileNotFoundError(resolved)

    size_mb = resolved.stat().st_size / (1024 * 1024)
    if size_mb > settings.max_file_size_mb:
        raise PathSecurityError(
            f"File {resolved} is {size_mb:.1f}MB, exceeds limit of {settings.max_file_size_mb}MB"
        )

    return resolved


def _is_blocked_address(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return True

    for _family, _type, _proto, _canonname, sockaddr in infos:
        ip = ipaddress.ip_address(sockaddr[0])
        if (
            ip.is_private
            or ip.is_loopback
            or ip.is_link_local
            or ip.is_reserved
            or ip.is_multicast
            or ip.is_unspecified
        ):
            return True
    return False


def validate_url(url: str) -> str:
    """Reject non-http(s) schemes and any URL whose host resolves to a
    private/loopback/link-local/reserved address (SSRF guard)."""
    parts = urlsplit(url)

    if parts.scheme not in ALLOWED_URL_SCHEMES:
        raise UrlSecurityError(f"Unsupported URL scheme: {parts.scheme!r}")

    if not parts.hostname:
        raise UrlSecurityError(f"URL has no host: {url}")

    if _is_blocked_address(parts.hostname):
        raise UrlSecurityError(f"URL host resolves to a disallowed address: {parts.hostname}")

    return url


def load_pdf(path: Path) -> list[Document]:
    """One Document per page; pages with no extractable text (e.g. scanned
    images) are dropped rather than indexed as empty content."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    source_doc_id = make_source_doc_id(str(path))

    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={"source": str(path), "source_doc_id": source_doc_id, "page": i + 1},
            )
        )
    return docs


def load_text(path: Path) -> list[Document]:
    """Loads a .txt/.md file as a single Document; empty/whitespace-only
    files produce no Document."""
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": str(path), "source_doc_id": make_source_doc_id(str(path))},
        )
    ]


def load_docx(path: Path) -> list[Document]:
    """Loads a .docx file's paragraphs as a single Document."""
    import docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    if not text.strip():
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": str(path), "source_doc_id": make_source_doc_id(str(path))},
        )
    ]


def load_url(url: str, _redirects_followed: int = 0) -> list[Document]:
    """Fetches a URL and extracts its visible text. Each redirect hop is
    re-validated through validate_url (an attacker-controlled redirect
    target is exactly as dangerous as the original URL)."""
    from bs4 import BeautifulSoup

    settings = get_settings()
    validate_url(url)

    response = requests.get(url, timeout=settings.url_fetch_timeout_seconds, allow_redirects=False, stream=True)

    if response.is_redirect or response.is_permanent_redirect:
        if _redirects_followed >= MAX_REDIRECTS:
            raise UrlSecurityError(f"Too many redirects fetching {url}")
        location = response.headers.get("Location")
        if not location:
            raise UrlSecurityError(f"Redirect from {url} had no Location header")
        return load_url(location, _redirects_followed=_redirects_followed + 1)

    max_bytes = settings.max_url_fetch_size_mb * 1024 * 1024
    chunks: list[bytes] = []
    total = 0
    for chunk in response.iter_content(chunk_size=8192):
        total += len(chunk)
        if total > max_bytes:
            raise UrlSecurityError(f"Response from {url} exceeded {settings.max_url_fetch_size_mb}MB limit")
        chunks.append(chunk)

    html = b"".join(chunks).decode(response.encoding or "utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.string.strip() if soup.title and soup.title.string else url
    text = soup.get_text(separator="\n", strip=True)

    return [
        Document(
            page_content=text,
            metadata={"source": url, "source_doc_id": make_source_doc_id(url), "title": title},
        )
    ]


def load_path(path: Path) -> list[Document]:
    """Routes an already-validated local path to the loader for its extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix in {".txt", ".md"}:
        return load_text(path)
    raise UnsupportedFileTypeError(f"Unsupported file extension: {suffix}")


def dispatch(source: str) -> list[Document]:
    """Routes a user-supplied source string to the right loader, applying
    the relevant security gate first."""
    parts = urlsplit(source)
    if parts.scheme in ALLOWED_URL_SCHEMES:
        return load_url(source)

    settings = get_settings()
    resolved = validate_local_path(Path(source), settings.docs_dir)
    return load_path(resolved)
