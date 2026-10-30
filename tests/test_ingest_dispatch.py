import socket
from types import SimpleNamespace

import docx as docx_module
import pytest
import requests

from ask_the_docs.ingest import (
    discover_sources,
    dispatch,
    dispatch_all_under_docs_dir,
    load_docx,
    load_url,
    validate_source,
)

SENTINEL = [object()]


class FakeResponse:
    def __init__(self, content: bytes, encoding=None, is_redirect=False, headers=None):
        self._content = content
        self.encoding = encoding
        self.is_redirect = is_redirect
        self.is_permanent_redirect = False
        self.headers = headers or {}

    def iter_content(self, chunk_size):
        yield self._content


def _allow_public_host(monkeypatch):
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda host, port: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 0))],
    )


def _patch_url_settings(monkeypatch, max_url_fetch_size_mb: int = 10):
    monkeypatch.setattr(
        "ask_the_docs.ingest.get_settings",
        lambda: SimpleNamespace(
            max_url_fetch_size_mb=max_url_fetch_size_mb, url_fetch_timeout_seconds=10
        ),
    )


def test_load_docx_reads_paragraphs(tmp_path):
    docx_path = tmp_path / "paper.docx"
    document = docx_module.Document()
    document.add_paragraph("Hello docx world.")
    document.save(str(docx_path))

    docs = load_docx(docx_path)

    assert len(docs) == 1
    assert "Hello docx world." in docs[0].page_content
    assert docs[0].metadata["source"] == str(docx_path)
    assert "source_doc_id" in docs[0].metadata


def test_load_docx_returns_empty_list_for_blank_document(tmp_path):
    docx_path = tmp_path / "blank.docx"
    docx_module.Document().save(str(docx_path))

    assert load_docx(docx_path) == []


def test_load_docx_keeps_heading_text_with_its_section(tmp_path):
    docx_path = tmp_path / "policy.docx"
    document = docx_module.Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph("Customers may request a refund within 30 days.")
    document.save(str(docx_path))

    docs = load_docx(docx_path)

    assert any("Refund Policy" in doc.page_content for doc in docs)


def test_load_docx_separates_sections_by_heading(tmp_path):
    docx_path = tmp_path / "policy.docx"
    document = docx_module.Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph("Refunds within 30 days.")
    document.add_heading("Shipping Policy", level=1)
    document.add_paragraph("Ships within 5 business days.")
    document.save(str(docx_path))

    docs = load_docx(docx_path)

    refund_doc = next(doc for doc in docs if "Refunds within 30 days." in doc.page_content)
    shipping_doc = next(doc for doc in docs if "Ships within 5 business days." in doc.page_content)
    assert "Shipping Policy" not in refund_doc.page_content
    assert "Refund Policy" not in shipping_doc.page_content
    assert "Refund Policy" in refund_doc.page_content
    assert "Shipping Policy" in shipping_doc.page_content


def test_load_docx_without_headings_still_produces_one_document(tmp_path):
    docx_path = tmp_path / "plain.docx"
    document = docx_module.Document()
    document.add_paragraph("Just a plain paragraph.")
    document.add_paragraph("Another plain paragraph.")
    document.save(str(docx_path))

    docs = load_docx(docx_path)

    assert len(docs) == 1
    assert "Just a plain paragraph." in docs[0].page_content
    assert "Another plain paragraph." in docs[0].page_content


def test_load_docx_title_style_also_starts_a_new_section(tmp_path):
    # python-docx's add_heading(level=0) uses the "Title" style, not
    # "Heading N" -- must be recognized as a section boundary too, not
    # silently folded into whatever group precedes it.
    docx_path = tmp_path / "titled.docx"
    document = docx_module.Document()
    document.add_paragraph("Preamble paragraph.")
    document.add_heading("Document Title", level=0)
    document.add_paragraph("Body under the title.")
    document.save(str(docx_path))

    docs = load_docx(docx_path)

    title_doc = next(doc for doc in docs if "Document Title" in doc.page_content)
    assert "Preamble paragraph." not in title_doc.page_content
    assert "Body under the title." in title_doc.page_content


def test_load_url_fetches_and_parses_html(monkeypatch):
    _allow_public_host(monkeypatch)
    _patch_url_settings(monkeypatch)
    html = b"<html><head><title>Hi</title></head><body><p>Hello world</p></body></html>"
    monkeypatch.setattr(requests, "get", lambda url, timeout, allow_redirects, stream: FakeResponse(html))

    docs = load_url("http://example.com/page")

    assert len(docs) == 1
    assert "Hello world" in docs[0].page_content
    assert docs[0].metadata["source"] == "http://example.com/page"
    assert docs[0].metadata["title"] == "Hi"


def test_load_url_parses_pdf_by_content_type(monkeypatch):
    """arxiv-style URLs (e.g. https://arxiv.org/pdf/2305.10601) serve a raw PDF
    with no .pdf suffix in the URL -- Content-Type is the only reliable signal."""
    _allow_public_host(monkeypatch)
    _patch_url_settings(monkeypatch)

    class FakePage:
        def extract_text(self):
            return "The paper's actual text."

    class FakeReader:
        def __init__(self, stream):
            self.pages = [FakePage()]

    monkeypatch.setattr("pypdf.PdfReader", FakeReader)
    monkeypatch.setattr(
        requests,
        "get",
        lambda url, timeout, allow_redirects, stream: FakeResponse(
            b"%PDF-1.4 fake bytes", headers={"Content-Type": "application/pdf"}
        ),
    )

    docs = load_url("https://arxiv.org/pdf/2305.10601")

    assert len(docs) == 1
    assert docs[0].page_content == "The paper's actual text."
    assert docs[0].metadata["source"] == "https://arxiv.org/pdf/2305.10601"
    assert docs[0].metadata["page"] == 1


def test_load_url_follows_redirect_and_revalidates_target(monkeypatch):
    _allow_public_host(monkeypatch)
    _patch_url_settings(monkeypatch)
    responses = {
        "http://example.com/old": FakeResponse(
            b"", is_redirect=True, headers={"Location": "http://example.com/new"}
        ),
        "http://example.com/new": FakeResponse(b"<html><body>New page</body></html>"),
    }
    monkeypatch.setattr(
        requests, "get", lambda url, timeout, allow_redirects, stream: responses[url]
    )

    docs = load_url("http://example.com/old")

    assert "New page" in docs[0].page_content
    assert docs[0].metadata["source"] == "http://example.com/new"


def test_load_url_rejects_too_many_redirects(monkeypatch):
    _allow_public_host(monkeypatch)
    _patch_url_settings(monkeypatch)
    monkeypatch.setattr(
        requests,
        "get",
        lambda url, timeout, allow_redirects, stream: FakeResponse(
            b"", is_redirect=True, headers={"Location": "http://example.com/next"}
        ),
    )

    import pytest

    from ask_the_docs.ingest import UrlSecurityError

    with pytest.raises(UrlSecurityError):
        load_url("http://example.com/start")


def test_load_url_enforces_size_limit(monkeypatch):
    _allow_public_host(monkeypatch)
    _patch_url_settings(monkeypatch, max_url_fetch_size_mb=0)
    monkeypatch.setattr(
        requests,
        "get",
        lambda url, timeout, allow_redirects, stream: FakeResponse(b"more than zero bytes"),
    )

    import pytest

    from ask_the_docs.ingest import UrlSecurityError

    with pytest.raises(UrlSecurityError):
        load_url("http://example.com/big")


def test_dispatch_routes_http_url_to_load_url(monkeypatch):
    monkeypatch.setattr("ask_the_docs.ingest.load_url", lambda url: SENTINEL)
    assert dispatch("http://example.com/paper") == SENTINEL


def test_dispatch_routes_local_path_to_load_path(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "ask_the_docs.ingest.get_settings",
        lambda: SimpleNamespace(docs_dir=tmp_path, max_file_size_mb=25),
    )
    doc_path = tmp_path / "paper.pdf"
    doc_path.write_bytes(b"%PDF fake")
    monkeypatch.setattr("ask_the_docs.ingest.load_path", lambda path: SENTINEL)

    assert dispatch("paper.pdf") == SENTINEL


def test_validate_source_checks_url_without_fetching_it(monkeypatch):
    called = []
    monkeypatch.setattr("ask_the_docs.ingest.validate_url", lambda url: called.append(url))
    monkeypatch.setattr("ask_the_docs.ingest.load_url", lambda url: (_ for _ in ()).throw(AssertionError("must not fetch")))

    validate_source("http://example.com/paper")

    assert called == ["http://example.com/paper"]


def test_validate_source_checks_local_path_without_loading_it(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "ask_the_docs.ingest.get_settings",
        lambda: SimpleNamespace(docs_dir=tmp_path, max_file_size_mb=25),
    )
    doc_path = tmp_path / "paper.pdf"
    doc_path.write_bytes(b"%PDF fake")
    monkeypatch.setattr("ask_the_docs.ingest.load_path", lambda path: (_ for _ in ()).throw(AssertionError("must not load")))

    validate_source("paper.pdf")  # must not raise, must not load


def test_validate_source_raises_for_an_ssrf_url(monkeypatch):
    from ask_the_docs.ingest import UrlSecurityError

    monkeypatch.setattr(socket, "getaddrinfo", lambda host, port: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", 0))])

    with pytest.raises(UrlSecurityError):
        validate_source("http://169.254.169.254/latest/meta-data")


def test_validate_source_raises_for_a_missing_local_file(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "ask_the_docs.ingest.get_settings",
        lambda: SimpleNamespace(docs_dir=tmp_path, max_file_size_mb=25),
    )

    with pytest.raises(FileNotFoundError):
        validate_source("missing.pdf")


def test_dispatch_all_under_docs_dir_returns_docs_and_file_count(tmp_path, monkeypatch):
    (tmp_path / "a.txt").write_text("hello")
    (tmp_path / "b.txt").write_text("world")
    monkeypatch.setattr(
        "ask_the_docs.ingest.get_settings",
        lambda: SimpleNamespace(docs_dir=tmp_path, max_file_size_mb=25),
    )
    settings = SimpleNamespace(docs_dir=tmp_path)

    docs, file_count = dispatch_all_under_docs_dir(settings)

    assert file_count == 2
    assert len(docs) == 2


def test_discover_sources_finds_allowed_extensions_only(tmp_path):
    (tmp_path / "paper.pdf").write_bytes(b"x")
    (tmp_path / "notes.md").write_text("x")
    (tmp_path / "ignore.exe").write_bytes(b"x")
    sub = tmp_path / "sub"
    sub.mkdir()
    (sub / "nested.txt").write_text("x")

    found = discover_sources(tmp_path)

    assert found == sorted(
        [tmp_path / "paper.pdf", tmp_path / "notes.md", sub / "nested.txt"]
    )
